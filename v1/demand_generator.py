import argparse
import json
import os
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

import openpyxl

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

PROFILE_XLSX    = "School Profile 2025.xlsx"
LOCATION_XLSX   = "School Location 2025.xlsx"
PROFILE_SHEET   = "SchoolProfile 2025"
LOCATION_SHEET  = "SchoolLocations 2025"
EXE_NAME        = "create_new_demand_points_windows_x86-64.exe"
OSRM_PORT       = 5000
OSRM_IMAGE      = "ghcr.io/project-osrm/osrm-backend"
OSRM_CONTAINER  = "osrm_sydney"

INCLUDE_TYPES   = {"primary", "secondary", "combined"}
TRAVEL_FRACTION = 0.9
MIN_ENROLMENT   = 50

BBOX_DEFAULT    = [150.53487298, -34.17324163, 151.34302098, -33.3645353]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def pop_size_for_enrolment(n: int) -> int:
    if n >= 1500: return 200
    if n >= 800:  return 150
    if n >= 400:  return 100
    if n >= 150:  return 75
    return 50


def in_bbox(lat: float, lon: float, bbox: list) -> bool:
    min_lon, min_lat, max_lon, max_lat = bbox
    return min_lon <= lon <= max_lon and min_lat <= lat <= max_lat


def make_code(name: str, index: int) -> str:
    skip = {
        "THE", "OF", "AND", "ST", "PUBLIC", "SCHOOL", "HIGH", "COLLEGE",
        "CATHOLIC", "PRIMARY", "COMMUNITY", "CHRISTIAN", "ANGLICAN",
        "GRAMMAR", "CENTRAL", "WEST", "EAST", "NORTH", "SOUTH"
    }
    words = name.upper().replace("-", " ").split()
    letters = [w[0] for w in words if w not in skip and w.isalpha()]
    code = "".join(letters[:5]) or "SCH"
    return f"{code}{index}"



def load_schools(script_dir: Path, bbox: list, min_enrolment: int) -> list:
    profile_path  = script_dir / PROFILE_XLSX
    location_path = script_dir / LOCATION_XLSX

    for p in [profile_path, location_path]:
        if not p.exists():
            print(f"  ERROR: Cannot find required file: {p.name}")
            print(f"         Expected location: {p}")
            sys.exit(1)

    print(f"  Reading {PROFILE_XLSX}...")
    wb = openpyxl.load_workbook(profile_path, read_only=True)
    ws = wb[PROFILE_SHEET]
    row_iter = ws.iter_rows(values_only=True)
    headers  = [str(h) for h in next(row_iter)]
    idx      = {h: i for i, h in enumerate(headers)}

    for col in ["ACARA SML ID", "School Name", "School Type", "Total Enrolments"]:
        if col not in idx:
            print(f"  ERROR: Expected column '{col}' not found in {PROFILE_XLSX}")
            print(f"         Found: {headers}")
            sys.exit(1)

    profiles = {}
    for row in row_iter:
        sid   = row[idx["ACARA SML ID"]]
        stype = str(row[idx["School Type"]] or "").strip().lower()
        if stype not in INCLUDE_TYPES:
            continue
        try:
            enrolment = int(str(row[idx["Total Enrolments"]] or "0").replace(",", ""))
        except ValueError:
            enrolment = 0
        if enrolment < min_enrolment:
            continue
        profiles[sid] = {
            "name":      str(row[idx["School Name"]] or sid).strip(),
            "type":      stype,
            "enrolment": enrolment,
        }
    wb.close()
    print(f"  Loaded {len(profiles)} primary/secondary/combined schools (enrolment >= {min_enrolment}).")

    print(f"  Reading {LOCATION_XLSX}...")
    wb = openpyxl.load_workbook(location_path, read_only=True)
    ws = wb[LOCATION_SHEET]
    row_iter = ws.iter_rows(values_only=True)
    headers  = [str(h) for h in next(row_iter)]
    idx      = {h: i for i, h in enumerate(headers)}

    for col in ["ACARA SML ID", "Latitude", "Longitude"]:
        if col not in idx:
            print(f"  ERROR: Expected column '{col}' not found in {LOCATION_XLSX}")
            print(f"         Found: {headers}")
            sys.exit(1)

    locations = {}
    for row in row_iter:
        sid = row[idx["ACARA SML ID"]]
        try:
            lat = float(row[idx["Latitude"]])
            lon = float(row[idx["Longitude"]])
            locations[sid] = {"lat": lat, "lon": lon}
        except (TypeError, ValueError):
            pass
    wb.close()
    print(f"  Loaded {len(locations)} school locations.")

    matched = []
    for sid, profile in profiles.items():
        if sid not in locations:
            continue
        loc = locations[sid]
        if not in_bbox(loc["lat"], loc["lon"], bbox):
            continue
        matched.append({**profile, "lat": loc["lat"], "lon": loc["lon"]})

    matched.sort(key=lambda s: s["enrolment"], reverse=True)

    seen = set()
    for i, school in enumerate(matched):
        code = make_code(school["name"], i)
        base, suffix = code, 0
        while code in seen:
            suffix += 1
            code = f"{base}{suffix}"
        seen.add(code)
        school["code"] = code

    types = {}
    for s in matched:
        types[s["type"]] = types.get(s["type"], 0) + 1
    print(f"\n  Schools within bbox: {len(matched)}")
    for t, c in sorted(types.items()):
        print(f"    {t.capitalize()}: {c}")
    total_students = sum(s["enrolment"] for s in matched)
    print(f"  Total enrolments:         {total_students:,}")
    print(f"  Est. daily transit demand: {int(total_students * TRAVEL_FRACTION):,}")

    return matched


def inject_schools(config: dict, schools: list) -> dict:
    def extend(key, values):
        config[key] = config.get(key, []) + values

    extend("universities",      [s["code"] for s in schools])
    extend("univ_loc",          [[s["lon"], s["lat"]] for s in schools])
    extend("univ_merge_within", [0] * len(schools))
    extend("students",          [s["enrolment"] for s in schools])
    extend("perc_oncampus",     [0.0] * len(schools))
    extend("univ_pop_size",     [pop_size_for_enrolment(s["enrolment"]) for s in schools])
    # univ_perc_travel is a single shared [on_campus_frac, off_campus_frac] list, not per-school
    config["univ_perc_travel"] = [TRAVEL_FRACTION, TRAVEL_FRACTION]
    return config



def osrm_healthy() -> bool:

    try:
        url = f"http://localhost:{OSRM_PORT}/nearest/v1/driving/151.2093,-33.8688"
        with urllib.request.urlopen(url, timeout=3) as r:
            return r.status == 200
    except urllib.error.HTTPError as e:
        return e.code in (200, 400)
    except Exception:
        return False


def start_osrm(osrm_arg: str, script_dir: Path):
    if osrm_healthy():
        print("  OSRM already running on port 5000 — skipping Docker start.")
        return

    # Resolve path
    osrm = Path(osrm_arg)
    if not osrm.is_absolute():
        osrm = script_dir / osrm
        
    def osrm_sidecars_exist(p: Path) -> bool:
        """Return True if any p.* sidecar files exist (e.g. sydney.osrm.names)."""
        return any(p.parent.glob(p.name + ".*"))

    if osrm.suffix == ".pbf":
        osrm = osrm.with_suffix(".osrm")

    if not osrm.exists() and not osrm_sidecars_exist(osrm):
        candidates = list(osrm.parent.glob("*.osrm.*"))
        if candidates:
            osrm = Path(str(candidates[0]).split(".osrm.")[0] + ".osrm")
            print(f"  Auto-detected OSRM data: {osrm.name}")
        else:
            print(f"  ERROR: Cannot find OSRM data file: {osrm}")
            print("         Run the three Docker pre-processing commands first:")
            print(f"           docker run -t -v \"${{PWD}}:/data\" {OSRM_IMAGE} osrm-extract -p /opt/car.lua /data/sydney.pbf")
            print(f"           docker run -t -v \"${{PWD}}:/data\" {OSRM_IMAGE} osrm-partition /data/sydney.osrm")
            print(f"           docker run -t -v \"${{PWD}}:/data\" {OSRM_IMAGE} osrm-customize /data/sydney.osrm")
            sys.exit(1)

    osrm_dir  = str(osrm.parent).replace("\\", "/")
    osrm_file = osrm.name

    subprocess.run(["docker", "rm", "-f", OSRM_CONTAINER], capture_output=True)

    cmd = [
        "docker", "run", "-d",
        "--name", OSRM_CONTAINER,
        "-p", f"{OSRM_PORT}:{OSRM_PORT}",
        "-v", f"{osrm_dir}:/data",
        OSRM_IMAGE,
        "osrm-routed",
        "--algorithm", "mld",
        f"/data/{osrm_file}"
    ]

    print(f"  Running: {' '.join(cmd)}")
    result = subprocess.run(cmd, capture_output=True, text=True)

    if result.returncode != 0:
        print("  ERROR: Docker failed to start OSRM:")
        print(result.stderr)
        sys.exit(1)

    print("  Waiting for OSRM to be ready", end="", flush=True)
    for _ in range(90):
        time.sleep(1)
        print(".", end="", flush=True)
        if osrm_healthy():
            print(" ready!")
            return

    print()
    print("  ERROR: OSRM did not become healthy within 90 seconds.")
    print(f"         Check logs with: docker logs {OSRM_CONTAINER}")
    sys.exit(1)

def run_demand_adder(config_path: Path, script_dir: Path):
    exe = script_dir / EXE_NAME
    if not exe.exists():
        print(f"  ERROR: Cannot find {EXE_NAME}")
        print(f"         Expected location: {exe}")
        print("         Download it from https://github.com/rslurry/Demand-Adder/releases")
        sys.exit(1)

    cmd = [str(exe), str(config_path)]
    print(f"  Running: {' '.join(cmd)}")
    result = subprocess.run(cmd, cwd=str(script_dir))

    if result.returncode != 0:
        print(f"\n  ERROR: Demand adder exited with code {result.returncode}")
        sys.exit(result.returncode)

def main():
    parser = argparse.ArgumentParser(
        description="Sydney demand pipeline: schools + OSRM + demand adder."
    )
    parser.add_argument(
        "--config", required=True,
        help="Path to your config JSON (e.g. sydney_config.json)"
    )
    parser.add_argument(
        "--pbf", required=True,
        help="Path to your processed OSRM data file (e.g. sydney.osrm)"
    )
    parser.add_argument(
        "--min-enrolment", type=int, default=MIN_ENROLMENT,
        help=f"Minimum school enrolment to include (default: {MIN_ENROLMENT})"
    )
    parser.add_argument(
        "--skip-schools", action="store_true",
        help="Skip school injection step"
    )
    parser.add_argument(
        "--skip-osrm", action="store_true",
        help="Skip OSRM Docker startup step"
    )
    args = parser.parse_args()

    script_dir  = Path(__file__).parent.resolve()
    config_path = Path(args.config)
    if not config_path.is_absolute():
        config_path = script_dir / config_path

    if not config_path.exists():
        print(f"ERROR: Config file not found: {config_path}")
        sys.exit(1)

    print("=" * 50)
    print("  Sydney Demand Pipeline")
    print("=" * 50)
    print(f"  Config:     {config_path.name}")
    print(f"  OSRM file:  {args.pbf}")
    print(f"  Script dir: {script_dir}")
    print()

    # Load config
    with open(config_path, encoding="utf-8") as f:
        config = json.load(f)

    bbox = config.get("bbox") or BBOX_DEFAULT

    # ---- Step 1: Schools ----
    if not args.skip_schools:
        print("[1/3] Injecting schools into config...")
        schools = load_schools(script_dir, bbox, args.min_enrolment)
        config  = inject_schools(config, schools)

        # Write updated config alongside original
        out_config = config_path.with_stem(config_path.stem + "_with_schools")
        with open(out_config, "w", encoding="utf-8") as f:
            json.dump(config, f, indent=2)
        print(f"\n  Saved updated config to: {out_config.name}")
        config_path = out_config
    else:
        print("[1/3] Skipping school injection.")

    # ---- Step 2: OSRM ----
    if not args.skip_osrm:
        print("\n[2/3] Starting OSRM via Docker...")
        start_osrm(args.pbf, script_dir)
    else:
        print("[2/3] Skipping OSRM startup.")

    # ---- Step 3: Demand adder ----
    print("\n[3/3] Running demand adder...")
    run_demand_adder(config_path, script_dir)

    print("\n" + "=" * 50)
    print("  Pipeline complete!")
    print(f"  Output: {config.get('output_demand_file', 'see config')}")
    print("=" * 50)


if __name__ == "__main__":
    main()